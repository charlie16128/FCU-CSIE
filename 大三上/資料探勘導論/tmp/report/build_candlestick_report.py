from __future__ import annotations

import argparse
import json
import math
import sys
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\逢甲大學\FCU Git\大三上\資料探勘導論")
PROJECT = ROOT / "candlestick_project"
RESULTS = PROJECT / "data" / "results"
ASSETS = ROOT / "tmp" / "report" / "assets"
OUTPUT = ROOT / "K線型態探索報告.docx"

ALL_PATTERNS = RESULTS / "all_patterns.csv"
TOP_BULLISH = RESULTS / "top10_bullish.csv"
TOP_BEARISH = RESULTS / "top10_bearish.csv"
FINAL_MODEL = RESULTS / "final_patterns.json"
TEST_OVERALL = RESULTS / "test_2026_overall_metrics.json"
TEST_SIGNALS = RESULTS / "test_2026_signals.csv"
TEST_PATTERN = RESULTS / "test_2026_pattern_metrics.csv"
TEST_STOCK = RESULTS / "test_2026_stock_metrics.csv"
VALIDATION = RESULTS / "validation_results.csv"

FLOWCHART = ASSETS / "program_flow.png"
BULLISH_FIGURE = ASSETS / "top10_bullish_report.png"
BEARISH_FIGURE = ASSETS / "top10_bearish_report.png"
TEST_FIGURE = ASSETS / "test_2026_report.png"

FONT_CJK = "Microsoft JhengHei"
FONT_LATIN = "Aptos"
FONT_MATH = "Cambria Math"
BLACK = "000000"
GREEN = "3F6F4E"
GREEN_LIGHT = "EAF2EC"
GRAY_LIGHT = "F4F4F4"
GRAY_BORDER = "D9D9D9"
GRAY_TEXT = "555555"


def load_inputs() -> dict[str, object]:
    required = (
        ALL_PATTERNS,
        TOP_BULLISH,
        TOP_BEARISH,
        FINAL_MODEL,
        TEST_OVERALL,
        TEST_SIGNALS,
        TEST_PATTERN,
        TEST_STOCK,
        VALIDATION,
    )
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing report inputs:\n" + "\n".join(missing))

    sys.path.insert(0, str(PROJECT))
    import config

    data: dict[str, object] = {
        "all_patterns": pd.read_csv(ALL_PATTERNS),
        "bullish": pd.read_csv(TOP_BULLISH),
        "bearish": pd.read_csv(TOP_BEARISH),
        "final_model": json.loads(FINAL_MODEL.read_text(encoding="utf-8")),
        "overall": json.loads(TEST_OVERALL.read_text(encoding="utf-8")),
        "signals": pd.read_csv(TEST_SIGNALS),
        "pattern_metrics": pd.read_csv(TEST_PATTERN),
        "stock_metrics": pd.read_csv(TEST_STOCK),
        "validation": pd.read_csv(VALIDATION),
        "stock_tickers": list(config.STOCK_TICKERS),
        "test_tickers": list(config.TEST_TICKERS),
    }
    return data


def validate_inputs(data: dict[str, object]) -> None:
    bullish = data["bullish"]
    bearish = data["bearish"]
    all_patterns = data["all_patterns"]
    final_model = data["final_model"]
    overall = data["overall"]
    signals = data["signals"]
    validation = data["validation"]
    assert isinstance(bullish, pd.DataFrame) and len(bullish) == 10
    assert isinstance(bearish, pd.DataFrame) and len(bearish) == 10
    assert isinstance(all_patterns, pd.DataFrame) and len(all_patterns) == 7661
    counts = all_patterns.groupby("direction").size().to_dict()
    assert counts == {"bearish": 3342, "bullish": 4319}
    assert int(all_patterns["cluster_size"].sum()) == 9246
    assert isinstance(final_model, dict)
    assert final_model["weight_preset"] == "shape_first"
    assert math.isclose(float(final_model["cluster_distance_threshold"]), 1.0)
    assert math.isclose(float(final_model["similarity_threshold"]), 1.6)
    assert isinstance(overall, dict)
    assert int(overall["total_signals"]) == 2
    assert int(overall["successful_signals"]) == 0
    assert math.isclose(float(overall["overall_accuracy"]), 0.0)
    assert isinstance(signals, pd.DataFrame) and len(signals) == 2
    assert isinstance(validation, pd.DataFrame) and len(validation) == 204

    raw_files = sorted((PROJECT / "data" / "raw").glob("*.csv"))
    assert len(raw_files) == 50
    raw_rows = sum(len(pd.read_csv(path, usecols=["Date"])) for path in raw_files)
    assert raw_rows == 104542
    assert len(data["stock_tickers"]) == 50
    assert len(data["test_tickers"]) == 10


def configure_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [FONT_CJK, "Microsoft YaHei", "Arial Unicode MS", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )


def build_flowchart() -> None:
    labels = [
        "以 yfinance 下載 50 檔股票資料",
        "清理 OHLCV 並排除除權息鄰近日期",
        "計算 10 維 K 線特徵與三日後報酬",
        "2018 至 2023 年探索候選型態",
        "看漲與看跌分開進行階層式聚類",
        "2024 至 2025 年搜尋權重與門檻",
        "依出現次數 準確率與報酬選出 Top 10",
        "鎖定 scaler 權重 門檻與型態",
        "2026 年固定 10 檔股票樣本外測試與 GUI",
    ]
    fig, ax = plt.subplots(figsize=(8, 10.2), dpi=220)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ys = np.linspace(0.93, 0.07, len(labels))
    for index, (label, y) in enumerate(zip(labels, ys, strict=True)):
        box = matplotlib.patches.FancyBboxPatch(
            (0.14, y - 0.034),
            0.72,
            0.068,
            boxstyle="round,pad=0.012,rounding_size=0.012",
            linewidth=1.7,
            edgecolor="#3F6F4E",
            facecolor="#F8FBF8" if index % 2 == 0 else "#EAF2EC",
        )
        ax.add_patch(box)
        ax.text(0.5, y, label, ha="center", va="center", fontsize=13, color="#111111")
        if index < len(labels) - 1:
            ax.annotate(
                "",
                xy=(0.5, ys[index + 1] + 0.043),
                xytext=(0.5, y - 0.043),
                arrowprops=dict(arrowstyle="-|>", color="#3F6F4E", lw=1.6),
            )
    ax.set_title("K 線型態探索程式流程", fontsize=20, color="#111111", pad=18, weight="bold")
    fig.savefig(FLOWCHART, bbox_inches="tight")
    plt.close(fig)


def decode_vector(value: object) -> np.ndarray:
    if isinstance(value, str):
        value = json.loads(value)
    vector = np.asarray(value, dtype=float)
    if vector.shape != (10,):
        raise ValueError("centroid_raw must contain ten values")
    return vector


def relative_candles(features: object) -> list[tuple[float, float, float, float]]:
    (
        upper,
        lower,
        body,
        prev_upper,
        prev_lower,
        prev_body,
        open_style,
        close_style,
        _volume_feature,
        _trend,
    ) = decode_vector(features)

    previous_close = 100.0
    previous_open = previous_close * (1.0 - prev_body / 100.0)
    previous_high = max(previous_open, previous_close) + prev_upper / 100.0 * previous_close
    previous_low = min(previous_open, previous_close) - prev_lower / 100.0 * previous_close

    denominator = 1.0 - close_style / 100.0
    current_close = previous_close / denominator if abs(denominator) > 1e-6 else previous_close
    current_open = previous_close + open_style / 100.0 * current_close
    body_open = current_close * (1.0 - body / 100.0)
    current_open = float(np.mean([current_open, body_open]))
    current_high = max(current_open, current_close) + upper / 100.0 * current_close
    current_low = min(current_open, current_close) - lower / 100.0 * current_close
    return [
        (previous_open, previous_high, previous_low, previous_close),
        (current_open, current_high, current_low, current_close),
    ]


def draw_candles(ax: plt.Axes, candles: list[tuple[float, float, float, float]]) -> None:
    for x, (open_price, high, low, close) in enumerate(candles):
        color = "#D62828" if close >= open_price else "#228B22"
        ax.vlines(x, low, high, color=color, linewidth=1.4)
        bottom = min(open_price, close)
        height = max(abs(close - open_price), 0.04)
        rect = matplotlib.patches.Rectangle(
            (x - 0.24, bottom),
            0.48,
            height,
            facecolor=color,
            edgecolor=color,
            linewidth=1.0,
        )
        ax.add_patch(rect)
    ax.set_xlim(-0.55, 1.55)
    ax.set_xticks([0, 1], ["前一日", "當日"], fontsize=8)
    ax.tick_params(axis="y", labelsize=7)
    ax.grid(axis="y", color="#E6E6E6", linewidth=0.6, alpha=0.8)
    for spine in ax.spines.values():
        spine.set_color("#AAAAAA")
        spine.set_linewidth(0.7)


def build_top10_figure(frame: pd.DataFrame, output: Path, title: str) -> None:
    fig, axes = plt.subplots(5, 2, figsize=(8, 10.2), dpi=220)
    for rank, (ax, (_, row)) in enumerate(zip(axes.ravel(), frame.iterrows(), strict=True), start=1):
        draw_candles(ax, relative_candles(row["centroid_raw"]))
        ax.set_title(
            f"{rank:02d}  {row['pattern_id']}\n"
            f"次數 {int(row['occurrence_count'])}   準確率 {float(row['accuracy']):.1%}   "
            f"平均報酬 {float(row['average_directional_profit']):.2%}",
            fontsize=9.2,
            color="#111111",
            pad=4,
        )
    fig.suptitle(title, fontsize=18, weight="bold", color="#111111", y=0.995)
    fig.tight_layout(rect=(0.02, 0.02, 0.98, 0.972), h_pad=1.35, w_pad=1.0)
    fig.savefig(output, bbox_inches="tight")
    plt.close(fig)


def build_test_figure(data: dict[str, object]) -> None:
    pattern = data["pattern_metrics"]
    stock = data["stock_metrics"]
    assert isinstance(pattern, pd.DataFrame)
    assert isinstance(stock, pd.DataFrame)
    pattern = pattern.loc[pattern["number_of_matches"].gt(0)].copy()
    stock = stock.loc[stock["number_of_signals"].gt(0)].copy()

    fig, axes = plt.subplots(1, 2, figsize=(8, 4.4), dpi=220)
    items = [
        (axes[0], pattern["pattern_id"], pattern["test_average_directional_profit"] * 100, "有訊號的型態"),
        (axes[1], stock["ticker"], stock["average_directional_profit"] * 100, "有訊號的股票"),
    ]
    for ax, labels, values, title in items:
        values_array = np.asarray(values, dtype=float)
        colors = ["#3F6F4E" if value >= 0 else "#B53A3A" for value in values_array]
        bars = ax.barh(list(labels), values_array, color=colors, height=0.55)
        ax.axvline(0, color="#666666", linewidth=0.8)
        ax.set_title(title, fontsize=13, weight="bold", color="#111111")
        ax.set_xlabel("平均方向性報酬百分比", fontsize=9)
        ax.grid(axis="x", color="#E6E6E6", linewidth=0.6)
        ax.tick_params(labelsize=9)
        for bar, value in zip(bars, values_array, strict=True):
            if value < -1.0:
                label_x = value / 2.0
                label_color = "FFFFFF"
                label_align = "center"
            else:
                label_x = value + 0.16
                label_color = "111111"
                label_align = "left"
            ax.text(
                label_x,
                bar.get_y() + bar.get_height() / 2,
                f"{value:.2f}%",
                va="center",
                ha=label_align,
                fontsize=9,
                color=f"#{label_color}",
            )
        span = max(1.0, float(np.max(np.abs(values_array))) * 1.25)
        ax.set_xlim(-span, span)
        for spine in ax.spines.values():
            spine.set_color("#BBBBBB")
    fig.suptitle("2026 年樣本外測試結果", fontsize=17, weight="bold", color="#111111", y=0.98)
    fig.text(
        0.5,
        0.015,
        "其餘 18 個型態與 8 檔測試股票沒有產生訊號",
        ha="center",
        fontsize=10,
        color="#555555",
    )
    fig.tight_layout(rect=(0.02, 0.06, 0.98, 0.92), w_pad=2.2)
    fig.savefig(TEST_FIGURE, bbox_inches="tight")
    plt.close(fig)


def build_assets(data: dict[str, object]) -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    configure_matplotlib()
    build_flowchart()
    build_top10_figure(data["bullish"], BULLISH_FIGURE, "看漲 Top 10 K 線型態")
    build_top10_figure(data["bearish"], BEARISH_FIGURE, "看跌 Top 10 K 線型態")
    build_test_figure(data)
    for path in (FLOWCHART, BULLISH_FIGURE, BEARISH_FIGURE, TEST_FIGURE):
        if not path.exists() or path.stat().st_size <= 20_000:
            raise RuntimeError(f"Report asset was not created correctly: {path}")


def set_style_font(style, name: str, size: float, *, bold: bool = False, color: str = BLACK) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN if name == FONT_CJK else name)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN if name == FONT_CJK else name)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)


def set_run_font(run, *, size: float | None = None, bold: bool | None = None, color: str = BLACK, name: str = FONT_CJK) -> None:
    run.font.name = name
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_LATIN if name == FONT_CJK else name)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN if name == FONT_CJK else name)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRAY_BORDER, size: str = "6") -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table) -> None:
    tblpr = table._tbl.tblPr
    layout = tblpr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")


def format_table(table, widths: list[float], *, center_columns: set[int] | None = None) -> None:
    center_columns = center_columns or set()
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout_fixed(table)
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            shade_cell(cell, GREEN if row_index == 0 else (GREEN_LIGHT if row_index % 2 == 0 else "FFFFFF"))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or column_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9.2,
                        bold=row_index == 0,
                        color="FFFFFF" if row_index == 0 else BLACK,
                    )


def set_keep_with_next(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        ppr.append(keep)


def set_keep_together(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    keep = ppr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        ppr.append(keep)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_keep_with_next(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=GRAY_TEXT)


def add_figure(doc: Document, path: Path, caption: str, width: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_together(paragraph)
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = doc.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    run = caption_paragraph.add_run(caption)
    set_run_font(run, size=9, color=GRAY_TEXT)


def add_math_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    math_para = OxmlElement("m:oMathPara")
    math_object = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math_object.append(math_run)
    math_para.append(math_object)
    paragraph._p.append(math_para)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_page_number(section) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=9, color=GRAY_TEXT)


def restart_page_numbers(section, start: int = 1) -> None:
    sectpr = section._sectPr
    pg_num = sectpr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sectpr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def configure_document(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)

    styles = doc.styles
    set_style_font(styles["Normal"], FONT_CJK, 11)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    styles["Normal"].paragraph_format.space_after = Pt(6)
    set_style_font(styles["Title"], FONT_CJK, 24, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(12)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    set_style_font(styles["Subtitle"], FONT_CJK, 13)
    styles["Subtitle"].font.italic = False
    styles["Subtitle"].paragraph_format.space_after = Pt(8)
    set_style_font(styles["Heading 1"], FONT_CJK, 16, bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(8)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    set_style_font(styles["Heading 2"], FONT_CJK, 13, bold=True)
    styles["Heading 2"].paragraph_format.space_before = Pt(7)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    set_style_font(styles["Caption"], FONT_CJK, 9, color=GRAY_TEXT)
    set_style_font(styles["List Bullet"], FONT_CJK, 11)
    set_style_font(styles["List Number"], FONT_CJK, 11)


def add_table_caption(doc: Document, number: int, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_with_next(paragraph)
    run = paragraph.add_run(f"表 {number}  {title}")
    set_run_font(run, size=10, bold=True)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("K 線型態探索與相似度比對報告")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("資料探勘導論作業一")
    doc.add_paragraph().paragraph_format.space_after = Pt(85)
    for label in ("姓名  ____________________", "學號  ____________________"):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(16)
        run = paragraph.add_run(label)
        set_run_font(run, size=13)
    doc.add_paragraph().paragraph_format.space_after = Pt(45)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date.add_run("2026 年 10 月")
    set_run_font(run, size=12, color=GRAY_TEXT)


def add_summary(doc: Document) -> None:
    heading = doc.add_heading("摘要", level=1)
    set_keep_with_next(heading)
    text = (
        "本研究在 2024 至 2025 年驗證資料中選出看漲與看跌各 10 組 K 線型態，但截至 "
        "2026 年 10 月 1 日的樣本外資料只出現 2 次看跌訊號，且兩次都未達到三個交易日後下跌超過 "
        "5% 的成功條件。因此，目前結果不足以證明此方法可以穩定協助投資人獲利。"
    )
    doc.add_paragraph(text)
    doc.add_paragraph(
        "方法以 50 檔臺灣股票的日 K 資料為基礎，將每個交易日轉換成 10 維數值特徵。"
        "2018 至 2023 年資料用於候選探索、標準化與聚類，2024 至 2025 年用於搜尋相似度權重與門檻，"
        "2026 年資料只用於鎖定模型後的樣本外測試。最終相似度使用標準化後的加權歐式距離。"
    )
    keywords = doc.add_paragraph()
    keywords.add_run("關鍵詞  ").bold = True
    keywords.add_run("K 線型態  特徵工程  相似度比對  階層式聚類  樣本外測試")


def add_dataset_section(doc: Document, data: dict[str, object]) -> None:
    doc.add_heading("資料集與資料處理", level=1)
    doc.add_paragraph(
        "本作業使用固定 50 檔臺灣股票，資料由 Yahoo Finance 透過 yfinance 取得。"
        "現有原始檔共 104,542 筆日資料，日期範圍為 2018 年 1 月 2 日至 2026 年 10 月 1 日。"
        "其中 7769.TW 上市時間較晚，因此可用資料為 466 筆；其餘多數股票各有 2,124 筆資料。"
    )
    doc.add_paragraph(
        "每筆資料至少包含 Open、High、Low、Close 與 Volume，並保留 Adj Close、Dividends 與 Stock Splits。"
        "程式移除重複日期、缺失或不合理價格，不以前值填補價格。遇到股利或股票分割時，"
        "將該日、前一交易日與後三個交易日標為不適合進行型態分析，避免把除權息跳空誤認為交易訊號。"
    )
    add_table_caption(doc, 1, "資料期間與用途")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "期間"
    table.rows[0].cells[1].text = "用途"
    table.rows[0].cells[2].text = "限制"
    rows = [
        ("2018 至 2023", "型態探索與聚類", "建立 scaler 及候選型態"),
        ("2024 至 2025", "驗證與參數搜尋", "選擇權重 門檻及 Top 10"),
        ("2026", "樣本外測試", "不得重新調整任何模型參數"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [1.35, 2.25, 2.85], center_columns={0})


def add_feature_section(doc: Document) -> None:
    doc.add_heading("K 線數值模型", level=1)
    doc.add_paragraph(
        "每個交易日以 10 個特徵描述。影線與實體保留價格方向，並加入前一日形狀、"
        "相對於前收盤價的位置、成交量及前五日趨勢，使模型同時比較形狀、位置、量能與短期方向。"
    )
    add_table_caption(doc, 2, "十維 K 線特徵")
    table = doc.add_table(rows=1, cols=3)
    headers = ("特徵", "公式", "說明")
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    features = [
        ("upper", "(H - max(O,C)) / C × 100", "當日上影線"),
        ("lower", "(min(O,C) - L) / C × 100", "當日下影線"),
        ("body", "(C - O) / C × 100", "當日實體 正值為上漲"),
        ("prev_upper", "upper[t-1]", "前一日上影線"),
        ("prev_lower", "lower[t-1]", "前一日下影線"),
        ("prev_body", "body[t-1]", "前一日實體"),
        ("open_style", "(O - C[t-1]) / C × 100", "開盤相對前收盤價的位置"),
        ("close_style", "(C - C[t-1]) / C × 100", "收盤相對前收盤價的位置"),
        ("volume_feature", "(V - 5MV) / V", "成交量相對五日均量"),
        ("trend", "(C[t-2] - C[t-7]) / C[t-7]", "前五個交易日趨勢"),
    ]
    for values in features:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [1.35, 3.05, 2.05], center_columns={0})
    doc.add_paragraph(
        "標籤使用三個交易日後報酬。若報酬大於 5%，列為看漲候選；若小於 -5%，列為看跌候選；"
        "其餘日期不作為候選型態。三日後報酬只作為標籤與評估結果，不放入 10 維輸入特徵。"
    )
    add_math_paragraph(doc, "return₃d = (C[t+3] - C[t]) / C[t]")


def add_algorithm_section(doc: Document) -> None:
    doc.add_heading("K 線型態相似度演算法", level=1)
    add_numbered(doc, "以 2018 至 2023 年有效資料計算各特徵的平均數與標準差，並做 Z-score 標準化。")
    add_math_paragraph(doc, "zᵢ = (xᵢ - μᵢ) / σᵢ")
    add_numbered(doc, "將三日後上漲超過 5% 與下跌超過 5% 的候選分開，以 Ward linkage 階層式聚類形成代表型態。")
    add_numbered(doc, "以群中心作為 Pattern，使用標準化後的加權歐式距離和交易日比對。")
    add_math_paragraph(doc, "D(A,B) = √(Σ wᵢ × (zAᵢ - zBᵢ)²)")
    add_numbered(doc, "若距離小於或等於門檻，判定該 Pattern 出現，再觀察三個交易日後是否符合正負 5% 的方向條件。")
    add_numbered(doc, "參數優先確保看漲與看跌方向都有足夠的合格型態，再比較出現次數、方向性報酬與準確率。")
    doc.add_paragraph(
        "程式比較 equal、shape_first 與 position_trend 三組可解釋權重，"
        "相似度門檻由 0.4 搜尋至 2.0，聚類距離門檻比較 1.0、1.25、1.5 與 2.0。"
        "型態至少要在驗證期出現 3 次且準確率達 55% 才能進入 Top 10；"
        "以 3 次為最低次數時，至少需要 2 次成功，仍可排除只出現一次的偶然結果。"
    )


def add_flow_section(doc: Document) -> None:
    doc.add_heading("程式流程", level=1)
    doc.add_paragraph(
        "程式把資料下載、資料清理、型態探索、驗證及最終測試分開。"
        "2026 年資料只進入最後一個測試階段，避免測試結果回頭影響模型。"
    )
    add_figure(doc, FLOWCHART, "圖 1  K 線型態探索程式流程", 5.8)


def add_parameter_table(doc: Document, data: dict[str, object]) -> None:
    final_model = data["final_model"]
    validation = data["validation"]
    selected = validation.loc[
        validation["weight_preset"].eq(final_model["weight_preset"])
        & validation["similarity_threshold"].eq(final_model["similarity_threshold"])
        & validation["cluster_distance_threshold"].eq(final_model["cluster_distance_threshold"])
    ].iloc[0]
    add_table_caption(doc, 3, "最終鎖定參數")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "項目"
    table.rows[0].cells[1].text = "結果"
    values = [
        ("Discovery scaler 樣本數", f"{int(final_model['scaler']['sample_count']):,}"),
        ("權重組合", "shape_first"),
        ("十維權重", "1.5, 1.5, 1.5, 1.0, 1.0,\n1.0, 1.25, 1.25, 0.75, 0.75"),
        ("聚類距離門檻", "1.0"),
        ("相似度門檻", "1.6"),
        ("合格看漲型態", f"{int(selected['qualified_bullish_patterns'])}"),
        ("合格看跌型態", f"{int(selected['qualified_bearish_patterns'])}"),
        ("Top 10 平均準確率", f"{float(selected['mean_accuracy']):.1%}"),
    ]
    for left, right in values:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    format_table(table, [2.65, 3.8], center_columns={1})


def add_validation_section(doc: Document, data: dict[str, object]) -> None:
    doc.add_heading("驗證結果與最佳型態", level=1)
    doc.add_paragraph(
        "2018 至 2023 年共有 9,246 筆大漲或大跌候選事件，其中看漲 5,200 筆、看跌 4,046 筆。"
        "在聚類距離門檻 1.0 下，候選形成 7,661 個代表型態，包括 4,319 個看漲型態與 3,342 個看跌型態。"
        "多數群只有一筆候選，代表嚴格距離下的型態非常分散，也增加驗證期覆蓋不足的風險。"
    )
    add_parameter_table(doc, data)
    doc.add_paragraph(
        "最終組合在驗證期得到 23 個合格看漲型態與 14 個合格看跌型態。"
        "兩個方向先依平均方向性報酬排序，再以加權距離 0.30 排除過度相似的型態，"
        "各保留 10 組作為鎖定模型。"
    )


def add_top10_table(doc: Document, frame: pd.DataFrame, number: int, title: str) -> None:
    add_table_caption(doc, number, title)
    table = doc.add_table(rows=1, cols=5)
    headers = ("排名", "型態編號", "出現次數", "準確率", "平均方向性報酬")
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for rank, (_, row) in enumerate(frame.iterrows(), start=1):
        cells = table.add_row().cells
        cells[0].text = str(rank)
        cells[1].text = str(row["pattern_id"])
        cells[2].text = str(int(row["occurrence_count"]))
        cells[3].text = f"{float(row['accuracy']):.1%}"
        cells[4].text = f"{float(row['average_directional_profit']):.2%}"
    format_table(table, [0.55, 1.25, 1.0, 1.05, 2.0], center_columns={0, 1, 2, 3, 4})


def add_top10_sections(doc: Document, data: dict[str, object]) -> None:
    bullish = data["bullish"]
    bearish = data["bearish"]
    doc.add_heading("看漲 Top 10", level=2)
    doc.add_paragraph(
        "看漲型態的驗證期平均方向性報酬介於 4.77% 至 10.63%。"
        "最高的 BULL-1618 出現 3 次，其中 2 次符合三日後上漲超過 5% 的條件。"
    )
    add_figure(doc, BULLISH_FIGURE, "圖 2  看漲 Top 10 K 線型態", 6.35)
    add_top10_table(doc, bullish, 4, "看漲 Top 10 驗證結果")

    doc.add_heading("看跌 Top 10", level=2)
    doc.add_paragraph(
        "看跌型態的驗證期平均方向性報酬介於 4.26% 至 12.20%。"
        "最高的 BEAR-3247 出現 5 次，其中 4 次符合三日後下跌超過 5% 的條件。"
    )
    add_figure(doc, BEARISH_FIGURE, "圖 3  看跌 Top 10 K 線型態", 6.35)
    add_top10_table(doc, bearish, 5, "看跌 Top 10 驗證結果")


def add_test_section(doc: Document, data: dict[str, object]) -> None:
    overall = data["overall"]
    signals = data["signals"]
    doc.add_heading("2026 樣本外測試", level=1)
    doc.add_paragraph(
        "測試股票在查看 2026 年結果前已固定為 2330.TW、2454.TW、2308.TW、2317.TW、3711.TW、"
        "2303.TW、3037.TW、2383.TW、2881.TW 與 2891.TW。現有資料截至 2026 年 10 月 1 日，"
        "因此結果尚未涵蓋完整年度。"
    )
    doc.add_paragraph(
        f"鎖定模型共產生 {int(overall['total_signals'])} 次訊號，兩次都是看跌訊號，"
        f"成功次數為 {int(overall['successful_signals'])}，整體準確率為 {float(overall['overall_accuracy']):.0%}，"
        f"平均方向性報酬為 {float(overall['average_directional_profit']):.2%}。"
        "看漲型態完全沒有訊號，顯示目前設定在樣本外資料中相當稀疏。"
    )
    add_table_caption(doc, 6, "2026 年實際訊號")
    table = doc.add_table(rows=1, cols=6)
    headers = ("型態", "股票", "日期", "三日報酬", "距離", "結果")
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for _, row in signals.iterrows():
        cells = table.add_row().cells
        values = (
            str(row["pattern_id"]),
            str(row["ticker"]),
            str(row["date"]),
            f"{float(row['return_3d']):.2%}",
            f"{float(row['distance']):.3f}",
            "未成功" if not bool(row["success"]) else "成功",
        )
        for index, value in enumerate(values):
            cells[index].text = value
    format_table(table, [1.0, 0.9, 1.25, 1.0, 0.75, 0.85], center_columns={0, 1, 2, 3, 4, 5})
    add_figure(doc, TEST_FIGURE, "圖 4  2026 年有訊號之型態與股票績效", 6.25)
    doc.add_paragraph(
        "BEAR-1581 在 2303.TW 的三日報酬為 -0.50%，方向雖然正確，但沒有達到作業規定的 -5%，因此判定失敗。"
        "BEAR-1056 在 3037.TW 出現後反而上漲 6.14%，方向性報酬為 -6.14%。"
    )


def add_discussion(doc: Document) -> None:
    doc.add_heading("是否可以幫助投資人獲利", level=1)
    doc.add_paragraph(
        "依目前樣本外測試，不能認定此方法可以幫助投資人穩定獲利。驗證期中的 Top 10 看起來具有較高平均方向性報酬，"
        "但 2026 年只出現 2 次訊號且全部失敗，訊號數也不足以估計可靠的成功率。"
    )
    add_bullet(doc, "歷史型態可能只符合 2018 至 2025 年的市場狀態，市場制度、波動度及資金偏好改變後，距離相近不代表報酬仍相同。")
    add_bullet(doc, "7,661 個聚類後型態相對分散，最終選擇仍可能受到多重比較與過度擬合影響。")
    add_bullet(doc, "最低出現次數為 3，雖可取得完整 Top 10，但估計值容易受少數事件影響。")
    add_bullet(doc, "模擬尚未計入手續費、交易稅、滑價與持有期間資金占用；看跌報酬也只是假設可放空的理論結果。")
    add_bullet(doc, "作業以三日後正負 5% 當作成功門檻，這是研究標籤，不一定是實務交易中風險報酬最佳的門檻。")
    doc.add_paragraph(
        "較合理的用途是把型態距離當作輔助篩選指標，再搭配風險管理、基本面或其他技術指標。"
        "若要繼續研究，應延長樣本外期間、增加跨市場資料、降低型態數量並使用交易成本後報酬評估。"
    )


def add_reflection(doc: Document) -> None:
    doc.add_heading("心得", level=1)
    doc.add_paragraph(
        "這次作業讓我理解，K 線圖雖然是視覺圖形，仍可拆成影線、實體、前一日位置、成交量與趨勢等數值特徵。"
        "只要先統一尺度，就能用距離計算比較兩個型態，而不必完全依靠人工辨識。"
    )
    doc.add_paragraph(
        "我認為最重要的部分不是找出報酬最高的圖形，而是把資料時間切分清楚。"
        "如果用 2026 年結果重新調整門檻或挑選型態，最後的準確率會失去樣本外測試的意義。"
        "本次結果也顯示，驗證期表現良好不代表下一段時間仍有效。"
    )
    doc.add_paragraph(
        "此外，型態出現次數與準確率需要一起考慮。只出現一兩次的型態即使報酬很高，也可能只是偶然。"
        "目前為了取得看漲與看跌各 10 組型態，最低次數設為 3，但 2026 年訊號仍然很少。"
        "未來若要提升可信度，我會優先增加資料量、簡化型態數量，並用更長的樣本外期間觀察。"
    )


def add_appendix(doc: Document, data: dict[str, object]) -> None:
    doc.add_heading("附錄", level=1)
    doc.add_heading("五十檔股票代碼", level=2)
    tickers = data["stock_tickers"]
    table = doc.add_table(rows=5, cols=10)
    for index, ticker in enumerate(tickers):
        row = index // 10
        column = index % 10
        table.cell(row, column).text = ticker
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell, 45, 35, 45, 35)
            shade_cell(cell, GREEN_LIGHT if row_index % 2 else "FFFFFF")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=7.5)

    doc.add_heading("固定測試股票", level=2)
    paragraph = doc.add_paragraph("  ".join(data["test_tickers"]))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("程式執行方式", level=2)
    doc.add_paragraph("在 candlestick_project 目錄依序執行下列命令：")
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(
        "python main.py --download    python main.py --analyze\n"
        "python main.py --test        python main.py --gui"
    )
    set_run_font(run, size=9.2, name="Consolas", color="333333")

    doc.add_heading("資料來源", level=2)
    doc.add_paragraph(
        "課程作業說明為根目錄的作業一找出K線型態 PDF。股票日資料由 Yahoo Finance 透過 yfinance 套件取得。"
        "本報告中的參數、Top 10 與 2026 測試數值均來自專案現有輸出。"
    )


def build_document(data: dict[str, object]) -> None:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "K 線型態探索與相似度比對報告"
    doc.core_properties.subject = "資料探勘導論作業一"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "K線 型態探索 相似度 資料探勘"

    add_cover(doc)
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    body_section.page_width = Inches(8.5)
    body_section.page_height = Inches(11)
    body_section.top_margin = Inches(0.8)
    body_section.bottom_margin = Inches(0.78)
    body_section.left_margin = Inches(0.82)
    body_section.right_margin = Inches(0.82)
    restart_page_numbers(body_section, 1)
    add_page_number(body_section)

    add_summary(doc)
    add_dataset_section(doc, data)
    doc.add_page_break()
    add_feature_section(doc)
    doc.add_page_break()
    add_algorithm_section(doc)
    add_flow_section(doc)
    doc.add_page_break()
    add_validation_section(doc, data)
    add_top10_sections(doc, data)
    add_test_section(doc, data)
    add_discussion(doc)
    add_reflection(doc)
    add_appendix(doc, data)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def validate_document() -> None:
    if not OUTPUT.exists() or OUTPUT.stat().st_size < 250_000:
        raise RuntimeError("DOCX is missing or unexpectedly small")
    doc = Document(OUTPUT)
    text_parts = [paragraph.text for paragraph in doc.paragraphs]
    text_parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    full_text = "\n".join(text_parts)
    for required in ("50 檔", "104,542", "7,661", "shape_first", "1.6", "2 次", "0%", "-2.82%"):
        if required not in full_text:
            raise AssertionError(f"Required report value is missing: {required}")
    if len(doc.inline_shapes) < 4:
        raise AssertionError("Expected at least four report figures")
    if len(doc.tables) < 7:
        raise AssertionError("Expected at least seven report tables")
    bullish_rows = next(table for table in doc.tables if table.cell(0, 1).text == "型態編號")
    if len(bullish_rows.rows) != 11:
        raise AssertionError("Top 10 table must contain ten data rows")
    with zipfile.ZipFile(OUTPUT) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_names = [name for name in archive.namelist() if name.startswith("word/footer") and name.endswith(".xml")]
        if not footer_names:
            raise AssertionError("Page-number footer is missing")
        footer_xml = "".join(archive.read(name).decode("utf-8") for name in footer_names)
        if "PAGE" not in footer_xml:
            raise AssertionError("Page-number field is missing")
        forbidden = ("turn0", "turn1", "tool citation", "placeholder")
        lowered = document_xml.lower()
        for token in forbidden:
            if token in lowered:
                raise AssertionError(f"Forbidden internal token found: {token}")
    print(f"Document validation passed: {OUTPUT} ({OUTPUT.stat().st_size:,} bytes)")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check-inputs", action="store_true")
    parser.add_argument("--assets-only", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    data = load_inputs()
    validate_inputs(data)
    if args.check_inputs:
        print("Input validation passed")
        return 0
    build_assets(data)
    if args.assets_only:
        print("Assets created")
        return 0
    build_document(data)
    validate_document()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
